"""
Testes do exportador Excel (output/exporta_excel.py) — modelo do pesquisador:

  - abas: Resumo, Gráficos, Dados (+ Qualidade, Reprodutibilidade quando há);
  - Dados: tabela estruturada ``TabDados`` (coluna "Período (UTC)" em faixa
    início–fim + "<COMPONENTE> (unidade)") e gráfico de linha;
  - Resumo: metadados, estatísticas por FÓRMULA em A16:G (AVERAGE/AVERAGEIF/
    MAX/MINIFS/SUM/kWh por dia) e gráfico comparativo no topo direito (F1);
  - Gráficos: energia diária por dia (SUMPRODUCT sobre o texto do período) +
    energia média diária por componente + 3 gráficos;
  - valores com 4 casas (precisão da fonte); kt fora das estatísticas.
"""

from __future__ import annotations

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from output.exporta_excel import exporta


def _df_uma_fonte(n_horas: int = 72) -> pd.DataFrame:
    ts = pd.date_range("2026-01-01 00:00", periods=n_horas, freq="h")
    h = ts.hour.values.astype(float)
    ghi = np.clip(np.cos((h - 15) / 5.0 * (np.pi / 2)), 0, None) * 600.0
    ghi = np.where((h >= 10) & (h <= 20), ghi, 0.0)
    return pd.DataFrame(
        {
            "timestamp": ts,
            "GHI": ghi.round(1),
            "DNI": (ghi * 1.3).round(1),
            "DHI": (ghi * 0.2).round(1),
            "BHI": (ghi * 0.8).round(1),
        }
    )


def _df_duas_fontes() -> pd.DataFrame:
    df = _df_uma_fonte()
    return pd.DataFrame(
        {
            "timestamp": df["timestamp"],
            "GHI_McClear": df["GHI"],
            "GHI_NASA": df["GHI"] * 0.7,
            "kt": np.where(df["GHI"] > 0, 0.7, np.nan),
        }
    )


def _metadados() -> dict:
    return {
        "local": "Botucatu",
        "latitude": -22.8867,
        "longitude": -48.4450,
        "altitude": 786.0,
        "periodo": "01/01/2026 a 03/01/2026",
        "fontes": "CAMS McClear",
        "passo_temporal": "1 hora",
        "email_soda": "pesquisador@unesp.br",
    }


def test_estrutura_basica_do_modelo(tmp_path):
    caminho = tmp_path / "saida.xlsx"
    exporta(_df_uma_fonte(), _metadados(), caminho)

    wb = load_workbook(caminho)
    # Ordem do modelo: Resumo -> Gráficos -> Dados.
    assert wb.sheetnames == ["Resumo", "Gráficos", "Dados"]

    ws = wb["Dados"]
    assert ws.freeze_panes == "A2"
    tab = ws.tables["TabDados"]
    assert tab.ref == "A1:E73"
    assert [c.name for c in tab.tableColumns] == [
        "Período (UTC)", "GHI (W/m²)", "DNI (W/m²)", "DHI (W/m²)", "BHI (W/m²)",
    ]
    # A 1ª coluna mostra a FAIXA início–fim (UTC), espelhando o site.
    assert ws["A2"].value == "01/01/2026 00:00–01:00"


def test_resumo_estatisticas_em_a16(tmp_path):
    caminho = tmp_path / "saida.xlsx"
    exporta(_df_uma_fonte(), _metadados(), caminho)

    r = load_workbook(caminho)["Resumo"]
    assert r["A1"].value == "PIRAXIS — Relatório de Radiação Solar"
    assert r["A2"].value == '=B4&" • "&B8'
    assert r["A3"].value == "Informações Gerais"
    rotulos = [r.cell(row=i, column=1).value for i in range(4, 12)]
    assert rotulos == [
        "Local", "Latitude", "Longitude", "Altitude (m)", "Período",
        "Fontes usadas", "Passo temporal", "Data de geração",
    ]
    assert r["B11"].number_format == "@"
    # Regressão (auditoria 2026-06-22): o e-mail SoDa é credencial pessoal e
    # NÃO entra no Excel — mesmo que venha nos metadados, o exportador ignora.
    valores_resumo = [
        str(c.value) for linha in r.iter_rows() for c in linha if c.value
    ]
    assert not any("pesquisador@unesp.br" in v for v in valores_resumo)
    assert not any("E-mail" in v for v in valores_resumo)

    # Estatísticas agora em A14 (título) / A16 (cabeçalho) / A17+ (dados).
    assert r["A14"].value == "Estatísticas de Radiação (W/m²)"
    assert [r.cell(row=16, column=j).value for j in range(1, 8)] == [
        "Componente", "Média", "Média Diurna", "Máximo", "Mínimo Diurno",
        "Energia (Wh/m²)", "kWh/m²/dia",
    ]
    assert r["A17"].value == "GHI"
    assert r["B17"].value == "=AVERAGE(TabDados[GHI (W/m²)])"
    assert r["C17"].value == '=AVERAGEIF(TabDados[GHI (W/m²)],">0")'
    assert r["D17"].value == "=MAX(TabDados[GHI (W/m²)])"
    assert r["E17"].value == (
        '=_xlfn.MINIFS(TabDados[GHI (W/m²)],TabDados[GHI (W/m²)],">0")'
    )
    assert r["F17"].value == "=SUM(TabDados[GHI (W/m²)])"
    assert r["G17"].value == "=F17/3/1000"  # 72 h = 3 dias
    assert r["B17"].number_format == "0.0000"


def test_aba_graficos(tmp_path):
    caminho = tmp_path / "saida.xlsx"
    exporta(_df_uma_fonte(), _metadados(), caminho)

    g = load_workbook(caminho)["Gráficos"]
    assert g["A1"].value == "Gráficos de Radiação Solar"
    assert g["A3"].value == "Energia diária por dia (kWh/m²)"
    assert g["A4"].value == "Data" and g["B4"].value == "GHI"
    assert g["A5"].value == "=DATE(2026,1,1)"
    # SUMPRODUCT que soma por dia a partir do TEXTO do período.
    assert "SUMPRODUCT" in g["B5"].value and "GHI (W/m²)" in g["B5"].value
    # Regressão (auditoria 2026-06-22): o token de ano em TEXT() tem de ser o
    # canônico "yyyy" — o .xlsx guarda fórmulas na forma invariante en-US e
    # "aaaa" (código de EXIBIÇÃO pt-BR) zeraria a soma diária ao abrir.
    assert 'TEXT($A5,"dd/mm/yyyy")' in g["B5"].value
    assert "aaaa" not in g["B5"].value
    # Energia média diária por componente referencia o Resumo.
    assert g["A10"].value == "Energia média diária por componente (kWh/m²/dia)"
    assert g["A12"].value == "=Resumo!A17"
    assert g["B12"].value == "=Resumo!G17"


def test_graficos_do_modelo(tmp_path):
    caminho = tmp_path / "saida.xlsx"
    exporta(_df_uma_fonte(), _metadados(), caminho)

    wb = load_workbook(caminho)
    tipos = [
        (aba.title, type(ch).__name__)
        for aba in wb.worksheets
        for ch in aba._charts
    ]
    # Resumo: barras (comparativo, topo direito). Gráficos: linha (perfil) +
    # 2 barras (energia diária e média). Dados: linha ao longo do tempo.
    assert tipos == [
        ("Resumo", "BarChart"),
        ("Gráficos", "LineChart"),
        ("Gráficos", "BarChart"),
        ("Gráficos", "BarChart"),
        ("Dados", "LineChart"),
    ]


def test_duas_fontes_kt_fora_das_estatisticas(tmp_path):
    caminho = tmp_path / "saida2.xlsx"
    exporta(_df_duas_fontes(), _metadados(), caminho)

    wb = load_workbook(caminho)
    assert wb.sheetnames == ["Resumo", "Gráficos", "Dados"]

    tab = wb["Dados"].tables["TabDados"]
    nomes = [c.name for c in tab.tableColumns]
    assert "kt" in nomes
    assert "GHI_McClear (W/m²)" in nomes

    # Estatísticas (A16 cabeçalho; A17+ componentes); kt não entra.
    r = wb["Resumo"]
    componentes = []
    i = 17
    while r.cell(row=i, column=2).value:  # enquanto houver fórmula de Média (col B)
        componentes.append(r.cell(row=i, column=1).value)
        i += 1
    assert componentes == ["GHI_McClear", "GHI_NASA"]
    assert "kt" not in componentes


def test_nan_vira_celula_vazia_e_sem_erros(tmp_path):
    df = _df_uma_fonte()
    df.loc[5, "GHI"] = np.nan
    caminho = tmp_path / "saida3.xlsx"
    exporta(df, _metadados(), caminho)

    wb = load_workbook(caminho)
    # Linha 7 (índice 5 + cabeçalho) com GHI vazio, não zero.
    assert wb["Dados"]["B7"].value is None
    # Nenhuma célula com erro de fórmula em nenhuma aba.
    for ws in wb.worksheets:
        for linha in ws.iter_rows():
            for cel in linha:
                if isinstance(cel.value, str):
                    assert not cel.value.startswith("#"), (
                        f"Erro em {ws.title}!{cel.coordinate}: {cel.value}"
                    )


def test_abas_qualidade_e_reprodutibilidade(tmp_path):
    """Com QC e reprodutibilidade, surgem as duas abas extras sem quebrar as outras."""
    from datetime import date

    from core.config import BOTUCATU
    from core.qualidade import analisar_qualidade
    from core.reprodutibilidade import gerar_reprodutibilidade

    df = _df_uma_fonte()
    qc = analisar_qualidade(df, BOTUCATU, "PT01H", ["CAMS McClear"])
    repro = gerar_reprodutibilidade(
        BOTUCATU, date(2026, 1, 1), date(2026, 1, 3), "PT01H", "1 hora",
        ["CAMS McClear"], ["GHI", "DNI", "DHI", "BHI"],
    )
    caminho = tmp_path / "saida_qc.xlsx"
    exporta(df, _metadados(), caminho, relatorio_qc=qc, reprodutibilidade=repro)

    wb = load_workbook(caminho)
    assert wb.sheetnames == [
        "Resumo", "Gráficos", "Dados", "Qualidade", "Reprodutibilidade",
    ]
    q_txt = " ".join(
        str(c.value) for row in wb["Qualidade"].iter_rows()
        for c in row if c.value is not None
    )
    assert "Controle de Qualidade" in q_txt
    assert "Completude" in q_txt
    r_txt = " ".join(
        str(c.value) for row in wb["Reprodutibilidade"].iter_rows()
        for c in row if c.value is not None
    )
    assert "Holmgren" in r_txt  # citação do pvlib
    assert "pesquisador@unesp.br" not in r_txt  # e-mail nunca na reprodutibilidade


def test_dados_exibidos_com_4_casas_sem_arredondar(tmp_path):
    """Valores de radiação exibidos com 4 casas (precisão da fonte), não 1."""
    caminho = tmp_path / "saida_fmt.xlsx"
    exporta(_df_uma_fonte(), _metadados(), caminho)
    wb = load_workbook(caminho)
    assert wb["Dados"]["B2"].number_format == "0.0000"   # célula de dado (GHI)
    assert wb["Resumo"]["B17"].number_format == "0.0000"  # Média do 1º componente


def test_passo_diario_usa_unidade_wh(tmp_path):
    ts = pd.date_range("2026-01-01", periods=10, freq="D")
    df = pd.DataFrame({"timestamp": ts, "GHI": np.full(10, 6500.0)})
    caminho = tmp_path / "saida4.xlsx"
    exporta(df, _metadados() | {"passo_temporal": "1 dia"}, caminho)

    wb = load_workbook(caminho)
    assert wb["Dados"]["B1"].value == "GHI (Wh/m²)"
    assert wb["Resumo"]["G17"].value == "=F17/10/1000"  # 10 dias
