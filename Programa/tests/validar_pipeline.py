"""
tests/validar_pipeline.py
=========================

Valida o pipeline interno extração -> Excel (sem interface) e os erros amigáveis.

Parte A (NÃO precisa de internet): pedir uma data recente demais (ontem) deve
produzir uma mensagem amigável de defasagem, sem traceback.

Parte B (precisa do e-mail SoDa + internet): gera o Excel do período padrão e
valida com pandas/openpyxl — abas esperadas, rótulos em português, 72 linhas,
sem células vazias nas colunas de radiação, e pico de GHI no Excel idêntico ao
do DataFrame de origem. Sem e-mail, a Parte B é pulada.

Uso:
    python tests/validar_pipeline.py [--email voce@dominio.com]
"""

from __future__ import annotations

import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

_RAIZ = Path(__file__).resolve().parent.parent
if str(_RAIZ) not in sys.path:
    sys.path.insert(0, str(_RAIZ))

import pandas as pd
from openpyxl import load_workbook

from tests._comum import (  # type: ignore
    APROVADO,
    PULADO,
    REPROVADO,
    PASSO_PADRAO,
    obter_email,
    periodo_padrao,
)

from core.pipeline import extrair_mcclear_para_excel

# Rótulos em português que devem aparecer na aba "Resumo".
ROTULOS_RESUMO = ["Local", "Período", "Fontes usadas", "Passo temporal"]
# Colunas da TabDados no formato do modelo oficial (passo horário -> W/m²).
COLUNAS_RADIACAO = ["GHI (W/m²)", "DNI (W/m²)", "DHI (W/m²)", "BNI (W/m²)"]


def _checar_erro_amigavel() -> tuple[bool, str]:
    """Parte A: data recente demais -> ValueError amigável (sem rede)."""
    amanha_demais = date.today() - timedelta(days=1)  # mais recente que hoje-2
    inicio = amanha_demais - timedelta(days=1)
    try:
        # E-mail com formato válido (a rejeição de data ocorre ANTES da API).
        extrair_mcclear_para_excel(
            "teste@exemplo.com", inicio, amanha_demais, PASSO_PADRAO
        )
    except ValueError as exc:
        msg = str(exc).lower()
        ok = ("defasagem" in msg or "disponível" in msg) and "traceback" not in msg
        return ok, f"mensagem: {str(exc)[:70]}…"
    except Exception as exc:  # noqa: BLE001
        return False, f"erro inesperado (não amigável): {type(exc).__name__}: {exc}"
    return False, "não levantou erro para data recente demais"


def _checar_excel(email: str) -> tuple[bool, str]:
    """Parte B: gera e valida o Excel do período padrão (precisa de internet)."""
    inicio, fim = periodo_padrao()
    with tempfile.TemporaryDirectory() as tmp:
        caminho = Path(tmp) / "validacao.xlsx"
        try:
            caminho, combinado = extrair_mcclear_para_excel(
                email, inicio, fim, PASSO_PADRAO, caminho_saida=caminho
            )
        except Exception as exc:  # noqa: BLE001
            return False, f"extração/exportação falhou: {exc}"

        problemas: list[str] = []
        wb = load_workbook(caminho)
        # O modelo oficial tem exatamente estas duas abas, nesta ordem.
        if wb.sheetnames != ["Resumo", "Dados"]:
            problemas.append(f"abas={wb.sheetnames} (esperado [Resumo, Dados])")

        # Tabela estruturada TabDados presente (as fórmulas dependem dela).
        if "TabDados" not in wb["Dados"].tables:
            problemas.append("tabela estruturada 'TabDados' ausente na aba Dados")

        # Rótulos em português na aba Resumo.
        resumo_txt = "\n".join(
            str(c.value)
            for row in wb["Resumo"].iter_rows()
            for c in row
            if c.value is not None
        )
        for rotulo in ROTULOS_RESUMO:
            if rotulo not in resumo_txt:
                problemas.append(f"rótulo ausente no Resumo: {rotulo!r}")

        # Aba Dados via pandas (só as colunas da TabDados: A até E).
        dados = pd.read_excel(caminho, sheet_name="Dados", usecols="A:E")
        if len(dados) != 72:
            problemas.append(f"linhas de dados={len(dados)} (esperado 72)")
        cols_rad = [c for c in COLUNAS_RADIACAO if c in dados.columns]
        if not cols_rad:
            problemas.append(
                f"nenhuma coluna do modelo encontrada em {list(dados.columns)}"
            )
        for c in cols_rad:
            if dados[c].isna().any():
                problemas.append(f"coluna de radiação com célula vazia: {c}")

        # Pico de GHI no Excel == no DataFrame de origem.
        col_ghi = "GHI (W/m²)"
        if col_ghi in dados.columns and "GHI" in combinado.columns:
            pico_excel = float(dados[col_ghi].max())
            pico_df = float(combinado["GHI"].max())
            if abs(pico_excel - pico_df) > 0.05:
                problemas.append(
                    f"pico GHI Excel={pico_excel:.2f} ≠ df={pico_df:.2f}"
                )

        if problemas:
            return False, "; ".join(problemas)
        return True, "Excel estruturalmente correto (abas, rótulos, 72 linhas, pico)"


def executar(email: str | None = None) -> tuple[str, str]:
    if email is None:
        email = obter_email()

    ok_a, det_a = _checar_erro_amigavel()
    print(f"      - {'OK ' if ok_a else 'FALHA'} erro amigável (data recente) · {det_a}")
    if not ok_a:
        return REPROVADO, "erro amigável não validado"

    if not email:
        return PULADO, "erro amigável OK; Excel pulado (sem e-mail SoDa)"

    ok_b, det_b = _checar_excel(email)
    print(f"      - {'OK ' if ok_b else 'FALHA'} Excel estrutural · {det_b}")
    if not ok_b:
        return REPROVADO, det_b
    return APROVADO, "erro amigável + Excel OK"


if __name__ == "__main__":
    status, detalhe = executar()
    print(f"\nvalidar_pipeline: {status} — {detalhe}")
    sys.exit(0 if status in (APROVADO, PULADO) else 1)
